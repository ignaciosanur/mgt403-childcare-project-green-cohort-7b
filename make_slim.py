"""
Build a slim NDCP file for PS4: IDs + median center-based weekly prices
(infant / toddler / preschool) and their source flags, 2015-2022.

Input : NDCP2015_2022.csv (output of trim_ndcp.py)
Output: NDCP2015_2022_center_median.csv / .xlsx

Columns kept (names as in the NDCP county data dictionary, Technical Report Appendix C):
  MCINFANT, MCTODDLER, MCPRESCHOOL = median full-time weekly price, center-based care,
  nominal dollars; *_flag = the NDCP's accompanying flag variable for each price.
No rows are dropped here (the metro-sample restriction happens in the analysis step).
"""
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

SRC = "NDCP2015_2022.csv"
OUT = "NDCP2015_2022_center_median"
COLS = ["STATE_NAME", "STATE_ABBREVIATION", "COUNTY_NAME", "COUNTY_FIPS_CODE", "STUDYYEAR",
        "MCINFANT", "MCInfant_flag", "MCTODDLER", "MCToddler_flag",
        "MCPRESCHOOL", "MCPreschool_flag"]

df = pd.read_csv(SRC, dtype={"COUNTY_FIPS_CODE": str}, usecols=COLS)[COLS]
df["COUNTY_FIPS_CODE"] = df["COUNTY_FIPS_CODE"].str.zfill(5)
for c in ["MCInfant_flag", "MCToddler_flag", "MCPreschool_flag"]:
    df[c] = df[c].astype("Int64")

df.to_csv(OUT + ".csv", index=False)

with pd.ExcelWriter(OUT + ".xlsx", engine="openpyxl") as xw:
    df.to_excel(xw, index=False, sheet_name="NDCP_center_median")
    ws = xw.sheets["NDCP_center_median"]
    ws.freeze_panes = "A2"
    for col_i, name in enumerate(COLS, start=1):
        letter = get_column_letter(col_i)
        ws.column_dimensions[letter].width = max(12, len(name) + 3)
        for cell in ws[letter]:
            cell.font = Font(name="Arial", bold=(cell.row == 1))
            if cell.row > 1 and name in ("MCINFANT", "MCTODDLER", "MCPRESCHOOL"):
                cell.number_format = "$#,##0.00"
            if cell.row > 1 and name == "COUNTY_FIPS_CODE":
                cell.number_format = "@"

print(df.shape)
print(df.groupby("STUDYYEAR")[["MCINFANT", "MCTODDLER", "MCPRESCHOOL"]].count())
